"""Synthetic regression tests; never reads actual user reports or survey data."""
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document

from build_evidence import evidence
from report_tools import inspect, verify
from insertion_checks import KEYS

ROOT = Path(__file__).resolve().parent


class SkillTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.cmp = pd.DataFrame({"domain": ["01", "02", "01", "02"], "year": [2019, 2019, 2020, 2020],
            "Direct": [.1, .2, .12, .22], "FH": [.11, .19, .13, .21], "MFH": [.12, .18, .14, .20],
            "Direct_MSE": [.02]*4, "FH_MSE": [.01, .02, .01, .02], "MFH_MSE": [.005, np.nan, .006, .02]})
        for m in ("FH", "MFH"):
            self.cmp[m+"_Bench"] = self.cmp[m]
            self.cmp[m+"_Bench_MSE"] = self.cmp[m+"_MSE"]
        self.sig = pd.DataFrame({"domain": ["01", "02", "01", "02"], "method": ["FH", "FH", "MFH", "MFH"],
            "diff": [.02, -.01, .03, -.02], "mse": [.001]*4, "lb": [-.02]*4, "ub": [.06]*4,
            "significant_unadjusted": [True, None, True, False], "significant_bh": [False, None, False, False],
            "significant_bonferroni": [None]*4})

    def workbook(self):
        p = self.root / "results.xlsx"
        with pd.ExcelWriter(p) as w:
            self.cmp.to_excel(w, sheet_name="Estimate comparison", index=False)
            self.sig.to_excel(w, sheet_name="Statistical significance", index=False)
        return p

    def test_unknown_flags_and_pairs(self):
        e = evidence(self.workbook(), "poverty", "proportion", "on")
        self.assertEqual(e["changes"][0]["flags"]["significant_bonferroni"]["significant"], None)
        self.assertEqual(e["paired_changes"][0]["jointly_tested"], 1)
        self.assertEqual(e["paired_changes"][0]["agreement"], 1)
        self.assertEqual(e["paired_levels"][0]["eligible"], 1)
        self.assertAlmostEqual(e["paired_levels"][0]["UFH_mse"]["mean"], .01)
        self.assertAlmostEqual(e["changes"][0]["width"]["mean"], 8)
        self.assertAlmostEqual(e["changes"][0]["mse"]["mean"], 10)

    def test_benchmark_not_inferred_from_equality(self):
        p = self.workbook()
        for setting in ("on", "unknown"):
            e = evidence(p, "poverty", "proportion", setting)
            self.assertEqual(e["settings"]["benchmark"], setting)
            self.assertEqual(len(e["benchmark_shifts"]), 4)
        e = evidence(p, "poverty", "proportion", "off")
        self.assertEqual(e["benchmark_shifts"], [])
        self.assertEqual(len(e["levels"]), 6)

    def test_duplicate_keys_rejected(self):
        self.cmp = pd.concat([self.cmp, self.cmp.iloc[:1]])
        with self.assertRaisesRegex(ValueError, "Duplicate"):
            evidence(self.workbook(), "poverty", "proportion", "on")

    def test_boolean_strings_and_bad_flags(self):
        self.sig["significant_unadjusted"] = ["TRUE", None, "FALSE", "FALSE"]
        e = evidence(self.workbook(), "poverty", "proportion", "off")
        self.assertEqual(e["paired_changes"][0]["UFH_only"], 1)
        self.sig.loc[0, "significant_unadjusted"] = "unknown"
        with self.assertRaisesRegex(ValueError, "Invalid significance"):
            evidence(self.workbook(), "poverty", "proportion", "off")

    def test_unit_modes_and_missing_columns(self):
        self.cmp = self.cmp.drop(columns=[c for c in self.cmp if "Bench" in c])
        p = self.workbook()
        for unit in ("percent", "currency", "log"):
            e = evidence(p, "mean_welfare", unit, "on")
            self.assertAlmostEqual(e["changes"][0]["width"]["mean"], .08)
            self.assertIsNone(e["benchmark_shifts"][0]["signed"]["mean"])
        with self.assertRaisesRegex(ValueError, "Mean welfare"):
            evidence(p, "mean_welfare", "proportion", "on")

    def test_formula_cells_rejected(self):
        self.cmp["FH"] = self.cmp.FH.astype(object)
        self.cmp.loc[0, "FH"] = "=1/10"
        with self.assertRaisesRegex(ValueError, "Formula cells"):
            evidence(self.workbook(), "poverty", "proportion", "off")

    def test_invalid_inference_not_rehabilitated(self):
        self.sig["joint_moment_status"] = ["regular", "regular", "invalid", "regular"]
        e = evidence(self.workbook(), "poverty", "proportion", "off")
        m = next(x for x in e["changes"] if x["method"] == "MFH")
        self.assertEqual(m["reported_flags"]["significant_unadjusted"]["significant"], 1)
        self.assertEqual(m["flags"]["significant_unadjusted"]["significant"], 0)
        self.assertEqual(m["flags"]["significant_unadjusted"]["tested"], 1)
        self.assertEqual(m["invalid_inference"], 1)
        self.assertEqual(e["paired_changes"][0]["jointly_tested"], 0)

    def spec(self):
        spec = {k: "Synthetic interpretation. All original estimates are retained." for k in KEYS}
        spec["meta"] = {"provider": "OpenAI", "model": "not recorded", "generated_utc": "2026-10-30 12:00:00 UTC",
                        "language": "English", "sources": "synthetic test fixtures"}
        spec["numeric_qa"] = {k: "limited" for k in KEYS}
        p = self.root / "sections.json"
        p.write_text(json.dumps(spec), encoding="utf-8")
        return p

    def insert(self, kind, src, out):
        return subprocess.run([sys.executable, str(ROOT/f"sae_ai_insert_{kind}.py"), str(src),
                               str(self.spec()), str(out)], capture_output=True, text=True)

    def test_html_insertion_and_preservation(self):
        src, out = self.root/"report.html", self.root/"enhanced.html"
        src.write_text('''<html><head></head><body><script>var unchanged=1;</script>
<table><caption>Run Setup</caption><tr><td>AI interpretations included</td><td>no</td></tr></table>
<div id="comparison-of-ufh-and-mfh"><h1>Comparison of UFH and MFH</h1>
<div class="ai-status-box">AI interpretations were not requested or were intentionally excluded from this report.</div>
<div id="normality-diagnostics"><h2>Normality Diagnostics</h2><p>Original diagnostics</p></div>
<div id="rate-mse-rmse-and-cv-comparisons"><h2>Rate, MSE, RMSE, and CV Comparisons</h2><p>Original rates</p></div>
<div id="poverty-maps"><h2>Poverty Maps</h2><img src="data:image/png;base64,YWJj"/></div>
<div id="rmse-maps"><h2>RMSE Maps</h2></div>
<div id="distribution-of-95-confidence-interval-widths"><h2>Distribution of 95% confidence-interval widths</h2>
<div class="excel-output-links"><a href="EU_SAE_results.xlsx">Workbook</a></div><p>Change maps use gray.</p></div>
</div>
<div id="session-information" number="4"><h1><span class="header-section-number">4</span> Session Information</h1><p>Original session</p></div>
</body></html>''', encoding="utf-8")
        before = src.read_bytes()
        result = self.insert("html", src, out)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(verify(src, out)["preservation"], "passed")
        self.assertEqual(src.read_bytes(), before)
        self.assertNotEqual(self.insert("html", out, self.root/"duplicate.html").returncode, 0)
        self.assertNotEqual(self.insert("html", src, src).returncode, 0)
        damaged = out.read_text().replace("Original rates", "Changed rates")
        out.write_text(damaged, encoding="utf-8")
        with self.assertRaisesRegex(AssertionError, "Original text"):
            verify(src, out)

    def test_docx_insertion_and_preservation(self):
        src, out = self.root/"report.docx", self.root/"enhanced.docx"
        d = Document()
        d.add_heading("Comparison of UFH and MFH", 1)
        d.add_paragraph("AI interpretations were not requested or were intentionally excluded from this report.")
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text = "AI interpretations included"; t.cell(0, 1).text = "no"
        for heading in ("Normality Diagnostics", "Rate, MSE, RMSE, and CV Comparisons", "Poverty Maps",
                        "RMSE Maps", "Distribution of 95% confidence-interval widths", "Change Maps"):
            d.add_heading(heading, 2)
            d.add_paragraph("Original " + heading)
        d.add_heading("4 Session Information", 1)
        d.add_paragraph("Original session")
        d.save(src)
        before = src.read_bytes()
        result = self.insert("docx", src, out)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(verify(src, out)["preservation"], "passed")
        self.assertEqual(src.read_bytes(), before)
        self.assertNotEqual(self.insert("docx", out, self.root/"duplicate.docx").returncode, 0)
        self.assertNotEqual(self.insert("docx", src, out).returncode, 0)


if __name__ == "__main__": unittest.main(verbosity=2)
