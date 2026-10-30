"""Inspect report evidence and check content preservation after insertion."""
import argparse
import base64
import hashlib
import json
import re
import zipfile
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document

STATUS = "AI interpretations were not requested or were intentionally excluded from this report."


def sha(data): return hashlib.sha256(data).hexdigest()
def norm(s): return re.sub(r"\s+", " ", s).strip()


def inspect(path):
    path = Path(path)
    result = {"sha256": sha(path.read_bytes()), "text": [], "tables": [], "images": []}
    if path.suffix.lower() == ".docx":
        d = Document(path)
        result["text"] = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            prev = t._element.getprevious()
            caption = "" if prev is None else "".join(prev.itertext())
            result["tables"].append({"caption_hint": caption,
                                     "rows": [[c.text for c in r.cells] for r in t.rows]})
        with zipfile.ZipFile(path) as z:
            result["parts"] = {n: sha(z.read(n)) for n in z.namelist() if not n.endswith("/")}
            result["images"] = [result["parts"][n] for n in sorted(result["parts"]) if n.startswith("word/media/")]
    elif path.suffix.lower() in (".html", ".htm"):
        s = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
        result["scripts"] = [str(t) for t in s.find_all("script")]
        result["links"] = [a.get("href") for a in s.find_all("a", href=True)]
        for img in s.find_all("img"):
            src = img.get("src", "")
            result["images"].append(sha(base64.b64decode(src.split(",", 1)[1]))
                                    if src.startswith("data:") and ";base64," in src else src)
        for t in s.find_all("table"):
            caption = t.find("caption")
            result["tables"].append({"caption_hint": caption.get_text(" ", strip=True) if caption else "",
                "rows": [[c.get_text(" ", strip=True) for c in row.find_all(["td", "th"])] for row in t.find_all("tr")]})
        for t in s.find_all(["script", "style", "head"]): t.decompose()
        result["text"] = list(s.stripped_strings)
        result["ai_keys"] = [t.get("id") for t in s.select("section.ai-interpretation-block")]
    else:
        raise ValueError("Supply HTML or DOCX, not an unverified .dox file")
    return result


def verify(src, out):
    a, b = inspect(src), inspect(out)
    assert [t["rows"] for t in a["tables"]] == [t["rows"] for t in b["tables"]], "Original tables changed"
    assert a["images"] == b["images"], "Original images changed"
    if "parts" in a:
        assert set(a["parts"]) == set(b["parts"]), "DOCX package parts changed"
        assert all(v == b["parts"][k] for k, v in a["parts"].items() if k != "word/document.xml"), "Non-document DOCX parts changed"
    else:
        assert a["scripts"] == b["scripts"], "Original scripts changed"
        assert a["links"] == b["links"], "Original links changed"
        keys = ["overview", "normality", "rates", "precision", "poverty_maps", "change_significance", "change_maps", "recommendation"]
        assert sorted(b["ai_keys"]) == sorted("ai-"+k for k in keys), "Missing or duplicate AI blocks"
    # The added text may split HTML text nodes, so compare the original normalized
    # text stream in sequence, excluding only the obsolete AI-off placeholder.
    target = norm(" ".join(b["text"]))
    cursor = 0
    for text in a["text"]:
        text = norm(text).replace(STATUS, "").strip()
        if not text: continue
        position = target.find(text, cursor)
        assert position >= 0, "Original text changed or reordered: " + text[:120]
        cursor = position + len(text)
    assert "These AI interpretations were added after the run" in target, "Missing post-run provenance"
    if "parts" in a:
        assert target.count("(added after the run)") == 8, "Missing or duplicate Word AI footers"
    return {"preservation": "passed", "source_sha256": a["sha256"], "output_sha256": b["sha256"],
            "visual_qa": "not checked by this script", "statistical_qa": "not checked by this script"}


if __name__ == "__main__":
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("command", choices=["inspect", "verify"])
    p.add_argument("source"); p.add_argument("destination")
    args = p.parse_args()
    if args.command == "inspect":
        if Path(args.destination).exists(): p.error("Output exists; choose a new evidence path")
        Path(args.destination).write_text(json.dumps(inspect(args.source), ensure_ascii=False, indent=2), encoding="utf-8")
    else:
        print(json.dumps(verify(args.source, args.destination), indent=2))
